from __future__ import annotations

import io
import zipfile

from promptharness.core.documents import ReadError, ReadResult, register

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MAX_OOXML_UNCOMPRESSED_BYTES = 100 * 1024 * 1024


def check_ooxml_size(data: bytes, kind: str) -> None:
    """Reject a non-zip, or a zip whose declared uncompressed size is over the limit.

    Called before python-docx / python-pptx / openpyxl see the data, so a small file
    that would inflate to gigabytes (a zip bomb) is refused without decompressing it.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            total = sum(info.file_size for info in z.infolist())
    except (zipfile.BadZipFile, ValueError, OSError) as e:
        raise ReadError(f"not a valid {kind} file") from e
    if total > MAX_OOXML_UNCOMPRESSED_BYTES:
        raise ReadError("file is too large when decompressed")


@register(".pdf")
def read_pdf(data: bytes, name: str) -> ReadResult:
    if not data.startswith(b"%PDF"):
        raise ReadError("not a valid PDF")
    from pypdf import PdfReader
    from pypdf.errors import DependencyError

    reader = PdfReader(io.BytesIO(data))
    try:
        if reader.is_encrypted and not reader.decrypt(""):
            raise ReadError("password-protected")
        pages = list(reader.pages)
        texts = [(p.extract_text() or "").strip() for p in pages]
    except DependencyError as e:
        raise ReadError("encrypted PDF needs the 'cryptography' package") from e

    warnings = tuple(
        f"{name}: page {i} has no extractable text" for i, t in enumerate(texts, 1) if not t
    )
    if not any(texts):
        raise ReadError(
            "no extractable text; it may be a scan. Convert the pages to images and attach those"
        )
    text = "".join(f"\n\n--- page {i} ---\n{t}" for i, t in enumerate(texts, 1))
    return ReadResult(text=text.strip("\n"), mime="application/pdf", pages=len(pages),
                      warnings=warnings)


def _row_text(cells) -> str:
    return " | ".join(c.text.strip() for c in cells)


@register(".docx")
def read_docx(data: bytes, name: str) -> ReadResult:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ReadError("not a valid .docx file")
    check_ooxml_size(data, ".docx")
    import docx
    from docx.table import Table

    try:
        doc = docx.Document(io.BytesIO(data))
    except KeyError as e:  # a zip that is not a Word package
        raise ReadError("not a valid .docx file") from e
    lines: list[str] = []
    for block in doc.iter_inner_content():
        if isinstance(block, Table):
            lines.extend(_row_text(row.cells) for row in block.rows)
        elif block.text.strip():
            lines.append(block.text)
    return ReadResult(text="\n".join(lines), mime=DOCX_MIME)


def _walk_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes)
        else:
            yield shape


@register(".pptx")
def read_pptx(data: bytes, name: str) -> ReadResult:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ReadError("not a valid .pptx file")
    check_ooxml_size(data, ".pptx")
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
    except KeyError as e:
        raise ReadError("not a valid .pptx file") from e
    out: list[str] = []
    count = 0
    for n, slide in enumerate(prs.slides, 1):
        count = n
        out.append(f"--- slide {n} ---")
        for shape in _walk_shapes(slide.shapes):
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                out.extend(_row_text(row.cells) for row in shape.table.rows)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"Notes: {notes}")
    return ReadResult(text="\n".join(out), mime=PPTX_MIME, pages=count)
