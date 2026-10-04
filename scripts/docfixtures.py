"""Generate small PDF, Word and PowerPoint files for tests."""

from __future__ import annotations

import io


def make_pdf(pages: list[str]) -> bytes:
    """Hand-written minimal PDF: one Helvetica text line per page."""
    objs: list[bytes] = []
    n = len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    font_id = 3 + 2 * n
    for i, text in enumerate(pages):
        page_id = 3 + 2 * i
        esc = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({esc}) Tj ET".encode("latin-1") if text else b""
        objs.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Contents {page_id + 1} 0 R "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>"
            ).encode()
        )
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n" % (len(objs) + 1)
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


def make_docx(
    paragraphs: list[str],
    table: list[list[str]] | None = None,
    blocks: list[str | list[list[str]]] | None = None,
) -> bytes:
    """Build a .docx. `blocks` (str = paragraph, list of rows = table, in order) wins."""
    import docx

    if blocks is None:
        blocks = [*paragraphs, *([table] if table else [])]
    doc = docx.Document()
    for b in blocks:
        if isinstance(b, str):
            doc.add_paragraph(b)
        else:
            t = doc.add_table(rows=len(b), cols=len(b[0]))
            for r, row in enumerate(b):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pptx(slides: list[dict]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # title and content
        slide.shapes.title.text = s.get("title", "")
        slide.placeholders[1].text = s.get("body", "")
        if s.get("table"):
            rows = s["table"]
            shape = slide.shapes.add_table(
                len(rows), len(rows[0]), Inches(1), Inches(4.5), Inches(6), Inches(1.5)
            )
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    shape.table.cell(r, c).text = cell
        if s.get("group_text"):
            grp = slide.shapes.add_group_shape()
            grp.shapes.add_textbox(Inches(1), Inches(6.5), Inches(4), Inches(0.5)).text_frame.text = (
                s["group_text"]
            )
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_ODF_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"'
)


def _odf(body: str, mime: str) -> bytes:
    import zipfile

    content = (
        f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {_ODF_NS}>'
        f"<office:body>{body}</office:body></office:document-content>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", mime, compress_type=zipfile.ZIP_STORED)
        z.writestr("content.xml", content)
    return buf.getvalue()


def make_odt(paragraphs: list[str]) -> bytes:
    from xml.sax.saxutils import escape

    body = "".join(
        f"<text:p>{escape(p)}</text:p>" if not p.startswith("# ")
        else f'<text:h text:outline-level="1">{escape(p[2:])}</text:h>'
        for p in paragraphs
    )
    return _odf(
        f"<office:text>{body}</office:text>", "application/vnd.oasis.opendocument.text"
    )


def make_ods(sheets: dict[str, list[list]]) -> bytes:
    from xml.sax.saxutils import escape

    out = []
    for name, rows in sheets.items():
        trs = ""
        for row in rows:
            tcs = ""
            for v in row:
                if isinstance(v, (int, float)):
                    tcs += (
                        f'<table:table-cell office:value-type="float" '
                        f'office:value="{v}"><text:p>{v}</text:p></table:table-cell>'
                    )
                elif v in (None, ""):
                    tcs += "<table:table-cell/>"
                else:
                    tcs += (
                        '<table:table-cell office:value-type="string">'
                        f"<text:p>{escape(str(v))}</text:p></table:table-cell>"
                    )
            trs += f"<table:table-row>{tcs}</table:table-row>"
        out.append(f'<table:table table:name="{escape(name)}">{trs}</table:table>')
    return _odf(
        f"<office:spreadsheet>{''.join(out)}</office:spreadsheet>",
        "application/vnd.oasis.opendocument.spreadsheet",
    )


def make_odp(slides: list[list[str]]) -> bytes:
    from xml.sax.saxutils import escape

    pages = "".join(
        f'<draw:page draw:name="page{i}">'
        + "".join(f"<draw:frame><draw:text-box><text:p>{escape(t)}</text:p></draw:text-box></draw:frame>" for t in texts)
        + "</draw:page>"
        for i, texts in enumerate(slides, 1)
    )
    return _odf(
        f"<office:presentation>{pages}</office:presentation>",
        "application/vnd.oasis.opendocument.presentation",
    )


def make_png(lines: list[str]) -> bytes:
    """Small white PNG with the lines drawn in black (large enough for a vision model)."""
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.load_default(size=28)
    line_h = 40
    img = Image.new("RGB", (640, 24 + line_h * len(lines)), "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((20, 16 + i * line_h), line, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue()
